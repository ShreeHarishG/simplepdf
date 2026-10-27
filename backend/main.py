# backend/main.py
"""
Simple PDF Engine - Backend API
Built for TFRENZY Simple PDF

Architecture Mapping:
--------------------------------------------------------------------------------------
Feature / Tool             | Primary Python Library   | Underlying Dependency
--------------------------------------------------------------------------------------
Merge / Split / Rotate     | pypdf                    | Native Python
PDF Compression            | pypdf                    | Native Python
PDF to Image (Thumbnails)  | PyMuPDF (fitz)           | C Engine / MuPDF
PDF to Word                | pdf2docx                 | PyMuPDF + python-docx
PDF to Excel               | pdfplumber + pandas      | pdfminer.six + openpyxl
Word / PPT to PDF          | subprocess               | Headless LibreOffice / docx2pdf
OCR (Scanned PDFs)         | pytesseract + pdf2image  | Tesseract Engine / Poppler
Watermark / Page Numbers   | ReportLab + pypdf        | Native Python
Photo to PDF               | Pillow (PIL)             | Native Python
--------------------------------------------------------------------------------------

To install:
    pip install -r requirements.txt

To run:
    uvicorn main:app --reload --port 8000
"""

import os
import io
import uuid
import shutil
import tempfile
import subprocess
from xml.sax.saxutils import escape
from typing import List, Optional
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
import base64

# Primary Libraries
import pypdf
from PIL import Image

app = FastAPI(
    title="Simple PDF API - TFRENZY",
    description="Online PDF toolkit backend engine using pypdf, pdf2docx, PyMuPDF, pdfplumber, and ReportLab",
    version="2.0.0",
)

# CORS configuration for Next.js frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)

TEMP_DIR = tempfile.gettempdir()


def _temp_path(name: str) -> str:
    safe = os.path.basename(name or "file")
    return os.path.join(TEMP_DIR, f"{uuid.uuid4().hex[:10]}_{safe}")


def _pdf_from_story(story, out_pdf: str, landscape: bool = False) -> None:
    from reportlab.lib.pagesizes import letter, landscape as landscape_size
    from reportlab.platypus import SimpleDocTemplate

    pagesize = landscape_size(letter) if landscape else letter
    doc = SimpleDocTemplate(
        out_pdf,
        pagesize=pagesize,
        leftMargin=48,
        rightMargin=48,
        topMargin=48,
        bottomMargin=48,
    )
    doc.build(story or [_plain_paragraph("This file had no readable text.")])


def _plain_paragraph(text: str, style_name: str = "Normal"):
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph

    styles = getSampleStyleSheet()
    cleaned = escape((text or "").replace("\x00", "")).replace("\n", "<br/>")
    return Paragraph(cleaned or "&nbsp;", styles[style_name])


def _docx_image_blobs(document) -> list:
    blobs = []
    ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for blip in document.element.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
        embed = blip.get(f"{ns}embed")
        rel = document.part.rels.get(embed) if embed else None
        if rel is not None:
            blobs.append(rel.target_part.blob)
    return blobs


def _docx_bytes_to_pdf(content: bytes, out_pdf: str) -> None:
    from docx import Document
    from reportlab.platypus import Spacer, Table, TableStyle, Image as FlowImage
    from reportlab.lib.utils import ImageReader
    from reportlab.lib import colors

    document = Document(io.BytesIO(content))
    story = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if text:
            style = "Heading2" if para.style and str(para.style.name).lower().startswith("heading") else "Normal"
            story.append(_plain_paragraph(text, style))
            story.append(Spacer(1, 8))
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append([_plain_paragraph(cell.text or " ") for cell in row.cells])
        if rows:
            flow = Table(rows, repeatRows=1)
            flow.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d4d4d8")),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f4f4f5")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(flow)
            story.append(Spacer(1, 12))
    for blob in _docx_image_blobs(document):
        try:
            reader = ImageReader(io.BytesIO(blob))
            width, height = reader.getSize()
            scale = min(1.0, 460.0 / max(width, 1))
            story.append(FlowImage(reader, width=width * scale, height=height * scale))
            story.append(Spacer(1, 10))
        except Exception:
            continue
    _pdf_from_story(story, out_pdf)


def _xlsx_bytes_to_pdf(content: bytes, out_pdf: str) -> None:
    from openpyxl import load_workbook
    from reportlab.platypus import Spacer, Table, TableStyle
    from reportlab.lib import colors

    workbook = load_workbook(io.BytesIO(content), data_only=True)
    story = []
    for sheet in workbook.worksheets:
        story.append(_plain_paragraph(sheet.title or "Sheet", "Heading2"))
        story.append(Spacer(1, 8))
        rows = []
        for row in sheet.iter_rows(max_row=80, max_col=10, values_only=True):
            values = [("" if cell is None else str(cell)) for cell in row]
            if any(cell.strip() for cell in values):
                rows.append([_plain_paragraph(cell) for cell in values])
        if rows:
            flow = Table(rows, repeatRows=1)
            flow.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d4d4d8")),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ecfdf5")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ]
                )
            )
            story.append(flow)
            story.append(Spacer(1, 16))
        else:
            story.append(_plain_paragraph("This sheet is empty."))
    _pdf_from_story(story, out_pdf, landscape=True)


def _emu_pt(value) -> float:
    try:
        return float(value) * 72.0 / 914400.0
    except (TypeError, ValueError):
        return 0.0


def _pptx_bytes_to_pdf(content: bytes, out_pdf: str) -> None:
    """One PDF page per slide, with pictures and text left where they sit on the slide."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from reportlab.pdfgen import canvas as pdf_canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.lib.colors import Color
    from reportlab.lib.utils import ImageReader

    presentation = Presentation(io.BytesIO(content))
    page_w = max(_emu_pt(presentation.slide_width), 72)
    page_h = max(_emu_pt(presentation.slide_height), 72)
    pdf = pdf_canvas.Canvas(out_pdf, pagesize=(page_w, page_h))

    def paint(shape, origin_x: float = 0.0, origin_y: float = 0.0) -> None:
        try:
            kind = shape.shape_type
        except Exception:
            return
        if kind == MSO_SHAPE_TYPE.GROUP:
            group_x = origin_x + _emu_pt(getattr(shape, "left", 0))
            group_y = origin_y + _emu_pt(getattr(shape, "top", 0))
            for child in shape.shapes:
                paint(child, group_x, group_y)
            return

        x = origin_x + _emu_pt(getattr(shape, "left", 0))
        y = origin_y + _emu_pt(getattr(shape, "top", 0))
        width = max(_emu_pt(getattr(shape, "width", 0)), 1)
        height = max(_emu_pt(getattr(shape, "height", 0)), 1)

        if kind == MSO_SHAPE_TYPE.PICTURE:
            try:
                pdf.drawImage(
                    ImageReader(io.BytesIO(shape.image.blob)),
                    x,
                    page_h - y - height,
                    width=width,
                    height=height,
                    preserveAspectRatio=True,
                    anchor="c",
                    mask="auto",
                )
            except Exception:
                pass
            return

        try:
            fore = shape.fill.fore_color.rgb
            pdf.setFillColor(Color(fore[0] / 255, fore[1] / 255, fore[2] / 255))
            pdf.rect(x, page_h - y - height, width, height, fill=1, stroke=0)
        except Exception:
            pass

        lines = []
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                joined = "  ".join((cell.text or "").strip() for cell in row.cells)
                if joined.strip():
                    lines.append((joined.strip(), 12))
        elif getattr(shape, "has_text_frame", False):
            for para in shape.text_frame.paragraphs:
                line = "".join(run.text or "" for run in para.runs) or (para.text or "")
                if not line.strip():
                    continue
                size = 18
                if para.runs and para.runs[0].font.size:
                    size = max(8, min(float(para.runs[0].font.size.pt), 72))
                lines.append((line.strip(), size))

        cursor = page_h - y - 6
        for line, size in lines:
            pdf.setFillColor(Color(0.12, 0.12, 0.14))
            pdf.setFont("Helvetica", size)
            current = ""
            for word in line.split():
                trial = f"{current} {word}".strip()
                if current and pdfmetrics.stringWidth(trial, "Helvetica", size) > max(width - 8, 12):
                    pdf.drawString(x + 4, cursor - size, current)
                    cursor -= size + 3
                    current = word
                else:
                    current = trial
            if current:
                pdf.drawString(x + 4, cursor - size, current)
                cursor -= size + 4

    if len(presentation.slides) == 0:
        pdf.showPage()
    for slide in presentation.slides:
        pdf.setFillColor(Color(1, 1, 1))
        pdf.rect(0, 0, page_w, page_h, fill=1, stroke=0)
        for shape in slide.shapes:
            paint(shape)
        pdf.showPage()
    pdf.save()


def _docx_has_media(path: str) -> bool:
    import zipfile

    try:
        with zipfile.ZipFile(path) as archive:
            return any(name.startswith("word/media/") for name in archive.namelist())
    except Exception:
        return False


def _pdf_has_images(content: bytes) -> bool:
    import pymupdf

    document = pymupdf.open(stream=content, filetype="pdf")
    try:
        return any(page.get_images() for page in document)
    finally:
        document.close()


def _pdf_bytes_to_docx(content: bytes, out_docx: str) -> None:
    in_pdf = _temp_path("source.pdf")
    with open(in_pdf, "wb") as handle:
        handle.write(content)

    converted = False
    # 1. Primary engine: pdf2docx Converter (full layout, text, tables, styles, and embedded images)
    try:
        from pdf2docx import Converter

        converter = Converter(in_pdf)
        converter.convert(out_docx, start=0, end=None)
        converter.close()
        converted = os.path.isfile(out_docx) and os.path.getsize(out_docx) > 1500
        if converted and _pdf_has_images(content) and not _docx_has_media(out_docx):
            converted = False
    except Exception as e:
        print(f"[pdf2docx] Warning: Converter failed ({e}), falling back to PyMuPDF image+text extractor.")
        converted = False
    finally:
        if os.path.isfile(in_pdf):
            try:
                os.remove(in_pdf)
            except Exception:
                pass

    if converted:
        return

    # 2. Enhanced Secondary engine: PyMuPDF (fitz) + python-docx (extracts text, embedded images, and renders scanned/graphical pages)
    try:
        import fitz
        import docx
        from docx.shared import Inches, Pt
        from PIL import Image as PILImage

        document = docx.Document()
        doc = fitz.open(stream=content, filetype="pdf")

        for page_index, page in enumerate(doc):
            if page_index > 0:
                document.add_page_break()

            page_text = (page.get_text() or "").strip()
            images_list = page.get_images(full=True)
            has_visuals = bool(images_list or page.get_drawings())

            # Keep the whole page look (pictures, lines, layout), then the editable text.
            if has_visuals:
                pix = page.get_pixmap(dpi=140)
                document.add_picture(io.BytesIO(pix.tobytes("png")), width=Inches(6.5))

            # Case B: Standard page with text blocks and/or embedded images
            blocks = page.get_text("blocks")
            text_blocks = [b for b in blocks if b[6] == 0 and b[4].strip()]
            if text_blocks:
                for b in text_blocks:
                    block_text = b[4].strip()
                    if not block_text:
                        continue
                    lines = [line.strip() for line in block_text.split("\n") if line.strip()]
                    for line in lines:
                        p = document.add_paragraph(line)
                        p.paragraph_format.space_after = Pt(4)
                        p.paragraph_format.line_spacing = 1.15
            elif page_text:
                for line in page_text.splitlines():
                    if line.strip():
                        p = document.add_paragraph(line.strip())
                        p.paragraph_format.space_after = Pt(4)

            # The page picture already includes these images.
            if has_visuals:
                continue
            seen_xrefs = set()
            for img_info in images_list:
                xref = img_info[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                try:
                    base_img = doc.extract_image(xref)
                    if base_img and "image" in base_img:
                        raw_bytes = base_img["image"]
                        try:
                            pil_im = PILImage.open(io.BytesIO(raw_bytes))
                            w, h = pil_im.size
                            if w < 16 or h < 16:
                                continue
                            max_width_in = 6.0
                            if w < 300:
                                max_width_in = min(3.0, (w / 96.0))
                            document.add_picture(io.BytesIO(raw_bytes), width=Inches(max_width_in))
                        except Exception:
                            document.add_picture(io.BytesIO(raw_bytes), width=Inches(5.0))
                except Exception as img_err:
                    print(f"Error extracting image xref {xref}: {img_err}")
                    continue

        doc.close()
        document.save(out_docx)
        return
    except Exception as e:
        print(f"[PyMuPDF] Fallback failed: {e}, falling back to pypdf fallback.")

    # 3. Last-ditch engine: pypdf (extracts text AND images)
    import docx
    from docx.shared import Inches

    document = docx.Document()
    reader = pypdf.PdfReader(io.BytesIO(content))
    wrote = False
    for index, page in enumerate(reader.pages, start=1):
        if index > 1:
            document.add_page_break()
        text = (page.extract_text() or "").strip()
        if text:
            document.add_paragraph(text)
            wrote = True
        try:
            for img_obj in page.images:
                document.add_picture(io.BytesIO(img_obj.data), width=Inches(5.0))
                wrote = True
        except Exception:
            pass
    if not wrote:
        document.add_paragraph("This PDF had no readable text or images.")
    document.save(out_docx)


def _office_to_pdf_file(content: bytes, filename: str, pdf_standard: str = "pdf1.7") -> str:
    stem, ext = os.path.splitext(filename or "document")
    ext = (ext or ".bin").lower()
    in_path = _temp_path(f"in{ext}")
    out_pdf = _temp_path(f"{stem or 'document'}.pdf")
    with open(in_path, "wb") as handle:
        handle.write(content)

    libreoffice_cmd = shutil.which("soffice") or shutil.which("libreoffice")
    converted = False
    convert_to = "pdf"

    if libreoffice_cmd:
        try:
            subprocess.run(
                [libreoffice_cmd, "--headless", "--convert-to", convert_to, in_path, "--outdir", TEMP_DIR],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
            )
            guessed = os.path.join(TEMP_DIR, f"{os.path.splitext(os.path.basename(in_path))[0]}.pdf")
            if os.path.isfile(guessed):
                shutil.move(guessed, out_pdf)
                converted = True
        except Exception:
            converted = False

    if not converted and ext in (".doc", ".docx"):
        try:
            from docx2pdf import convert
            convert(in_path, out_pdf)
            converted = os.path.isfile(out_pdf)
        except Exception:
            pass

    if not converted:
        try:
            if ext in (".doc", ".docx"):
                _docx_bytes_to_pdf(content, out_pdf)
            elif ext in (".xls", ".xlsx"):
                _xlsx_bytes_to_pdf(content, out_pdf)
            elif ext in (".ppt", ".pptx"):
                _pptx_bytes_to_pdf(content, out_pdf)
            else:
                raise ValueError("unsupported office type")
            converted = os.path.isfile(out_pdf)
        except Exception as exc:
            raise HTTPException(
                status_code=500,
                detail=f"Could not convert that Office file to PDF: {exc}",
            ) from exc

    if not converted or not os.path.isfile(out_pdf):
        raise HTTPException(status_code=500, detail="Could not convert that Office file to PDF.")
    return out_pdf


@app.get("/")
def health_check():
    return {
        "status": "online",
        "service": "Simple PDF Backend Engine (TFRENZY)",
        "tool_matrix": {
            "merge_pdf": {"library": "pypdf", "dependency": "Native Python"},
            "compress_pdf": {"library": "pypdf", "dependency": "Native Python"},
            "pdf_to_image": {"library": "PyMuPDF (fitz)", "dependency": "C Engine / MuPDF"},
            "pdf_to_word": {"library": "pdf2docx", "dependency": "PyMuPDF + python-docx"},
            "pdf_to_excel": {"library": "pdfplumber", "dependency": "pdfminer.six + openpyxl"},
            "word_to_pdf": {"library": "subprocess", "dependency": "Headless LibreOffice"},
            "excel_to_pdf": {"library": "subprocess", "dependency": "Headless LibreOffice"},
            "ppt_to_pdf": {"library": "subprocess", "dependency": "Headless LibreOffice"},
            "ocr_pdf": {"library": "pytesseract", "dependency": "Tesseract Engine"},
            "watermark_pdf": {"library": "ReportLab + pypdf", "dependency": "Native Python"},
            "photo_to_pdf": {"library": "Pillow (PIL)", "dependency": "Native Python"},
            "split_pdf": {"library": "pypdf", "dependency": "Native Python"},
            "rotate_pdf": {"library": "pypdf", "dependency": "Native Python"},
            "organize_pdf": {"library": "pypdf", "dependency": "Native Python"},
            "extract_text": {"library": "pypdf", "dependency": "Native Python"},
            "page_numbers": {"library": "ReportLab + pypdf", "dependency": "Native Python"},
            "protect_pdf": {"library": "pypdf", "dependency": "AES-256"},
            "unlock_pdf": {"library": "pypdf", "dependency": "Native Python"},
        },
    }


# ==============================================================================
# 1. MERGE PDF (Library: pypdf | Native Python)
# ==============================================================================
@app.post("/api/merge-pdf")
async def merge_pdf(files: List[UploadFile] = File(...)):
    if len(files) < 2:
        raise HTTPException(status_code=400, detail="Provide at least 2 PDF files to merge.")

    writer = pypdf.PdfWriter()
    try:
        for file in files:
            content = await file.read()
            writer.append(io.BytesIO(content))

        out_path = os.path.join(TEMP_DIR, "simplepdf_merged.pdf")
        with open(out_path, "wb") as f_out:
            writer.write(f_out)
        writer.close()

        return FileResponse(
            out_path,
            media_type="application/pdf",
            filename="simplepdf_merged.pdf",
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Merge error: {str(e)}")


# ==============================================================================
# 2. PDF COMPRESSION (Library: pypdf | Native Python)
# ==============================================================================
@app.post("/api/compress-pdf")
async def compress_pdf(
    file: UploadFile = File(...),
    level: str = Form("recommended"),
):
    try:
        content = await file.read()
        reader = pypdf.PdfReader(io.BytesIO(content))
        writer = pypdf.PdfWriter()
        writer.append(reader)
        for page in writer.pages:
            page.compress_content_streams()
        if reader.metadata:
            writer.add_metadata(reader.metadata)

        stem = os.path.splitext(file.filename or "document")[0]
        out_path = _temp_path(f"compressed_{stem}.pdf")
        with open(out_path, "wb") as f_out:
            writer.write(f_out)

        return FileResponse(
            out_path,
            media_type="application/pdf",
            filename=f"compressed_{stem}.pdf",
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Compression error: {str(e)}")


# ==============================================================================
# 3. PDF TO WORD (Library: pdf2docx | PyMuPDF + python-docx)
# ==============================================================================
@app.post("/api/pdf-to-word")
async def pdf_to_word(file: UploadFile = File(...)):
    try:
        content = await file.read()
        stem = os.path.splitext(file.filename or "document")[0]
        out_docx = _temp_path(f"{stem}.docx")
        _pdf_bytes_to_docx(content, out_docx)
        return FileResponse(
            out_docx,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=f"{stem}.docx",
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF to Word error: {str(e)}")


# ==============================================================================
# 4. WORD TO PDF (Library: subprocess Headless LibreOffice / docx2pdf)
# ==============================================================================
@app.post("/api/word-to-pdf")
async def word_to_pdf(file: UploadFile = File(...), pdf_standard: str = Form("pdf1.7")):
    try:
        content = await file.read()
        out_pdf = _office_to_pdf_file(content, file.filename or "document.docx", pdf_standard)
        stem = os.path.splitext(file.filename or "document")[0]
        return FileResponse(out_pdf, media_type="application/pdf", filename=f"{stem}.pdf")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Word to PDF error: {str(e)}")


@app.post("/api/office-to-pdf")
async def office_to_pdf(
    file: UploadFile = File(...),
    pdf_standard: str = Form("pdf1.7"),
):
    try:
        content = await file.read()
        out_pdf = _office_to_pdf_file(content, file.filename or "document.bin", pdf_standard)
        stem = os.path.splitext(file.filename or "document")[0]
        return FileResponse(out_pdf, media_type="application/pdf", filename=f"{stem}.pdf")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Office to PDF error: {str(e)}")


@app.post("/api/excel-to-pdf")
async def excel_to_pdf(file: UploadFile = File(...), pdf_standard: str = Form("pdf1.7")):
    return await office_to_pdf(file, pdf_standard)


@app.post("/api/ppt-to-pdf")
async def ppt_to_pdf(file: UploadFile = File(...), pdf_standard: str = Form("pdf1.7")):
    return await office_to_pdf(file, pdf_standard)


# ==============================================================================
# 5. PHOTO TO PDF (Library: Pillow | Native Python)
# ==============================================================================
@app.post("/api/photo-to-pdf")
async def photo_to_pdf(
    files: List[UploadFile] = File(...),
    orientation: str = Form("portrait"),
    page_size: str = Form("a4"),
    margin: str = Form("small"),
):
    if not files:
        raise HTTPException(status_code=400, detail="Provide at least one image file.")

    try:
        sizes = {"a4": (1240, 1754), "letter": (1275, 1650)}
        margins = {"none": 0, "small": 48, "big": 96}
        images = []
        for file in files:
            content = await file.read()
            img = Image.open(io.BytesIO(content)).convert("RGB")
            if page_size != "fit":
                page_w, page_h = sizes.get(page_size, sizes["a4"])
                if orientation == "landscape":
                    page_w, page_h = page_h, page_w
                pad = margins.get(margin, 48)
                canvas = Image.new("RGB", (page_w, page_h), "white")
                img.thumbnail((max(1, page_w - pad * 2), max(1, page_h - pad * 2)), Image.Resampling.LANCZOS)
                canvas.paste(img, ((page_w - img.width) // 2, (page_h - img.height) // 2))
                img = canvas
            images.append(img)

        out_path = os.path.join(TEMP_DIR, "simplepdf_photos.pdf")
        images[0].save(
            out_path,
            save_all=True,
            append_images=images[1:] if len(images) > 1 else [],
            quality=95,
        )

        return FileResponse(
            out_path,
            media_type="application/pdf",
            filename="photos_converted.pdf",
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Photo to PDF error: {str(e)}")


# ==============================================================================
# 6. PDF TO IMAGE / THUMBNAILS (Library: PyMuPDF fitz | C Engine / MuPDF)
# ==============================================================================
@app.post("/api/pdf-to-image")
async def pdf_to_image(
    file: UploadFile = File(...),
    page_num: int = Form(-1),
    pages: str = Form("all"),
    image_format: str = Form("png"),
    dpi: int = Form(150),
):
    try:
        import zipfile
        import fitz  # PyMuPDF
        content = await file.read()
        doc = fitz.open(stream=content, filetype="pdf")
        fmt = "jpeg" if image_format.lower() in ("jpg", "jpeg") else "png"
        ext = "jpg" if fmt == "jpeg" else "png"
        safe_dpi = 72 if dpi <= 72 else 300 if dpi >= 300 else 150
        indices = [min(page_num, len(doc) - 1)] if page_num >= 0 else _parse_pages(pages, len(doc))
        stem = os.path.splitext(file.filename or "document")[0]
        if len(indices) == 1:
            pix = doc[indices[0]].get_pixmap(dpi=safe_dpi)
            out_path = os.path.join(TEMP_DIR, f"{stem}_page.{ext}")
            pix.save(out_path)
            doc.close()
            return FileResponse(out_path, media_type=f"image/{fmt}", filename=f"{stem}_page_{indices[0] + 1}.{ext}")
        out_path = os.path.join(TEMP_DIR, f"{stem}_images.zip")
        with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for index in indices:
                pix = doc[index].get_pixmap(dpi=safe_dpi)
                archive.writestr(f"{stem}_page_{index + 1}.{ext}", pix.tobytes(fmt))
        doc.close()
        return FileResponse(out_path, media_type="application/zip", filename=f"{stem}_images.zip")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF to Image error: {str(e)}")


@app.post("/api/pdf-thumbnails")
async def pdf_thumbnails(file: UploadFile = File(...)):
    try:
        content = await file.read()
        pages = []
        try:
            import fitz

            doc = fitz.open(stream=content, filetype="pdf")
            limit = min(len(doc), 60)
            for index in range(limit):
                pix = doc[index].get_pixmap(dpi=110)
                encoded = base64.b64encode(pix.tobytes("png")).decode("ascii")
                pages.append(
                    {
                        "index": index,
                        "width": pix.width,
                        "height": pix.height,
                        "image": f"data:image/png;base64,{encoded}",
                    }
                )
            total = len(doc)
            doc.close()
            return JSONResponse({"pageCount": total, "truncated": total > limit, "pages": pages})
        except ImportError:
            reader = pypdf.PdfReader(io.BytesIO(content))
            total = len(reader.pages)
            limit = min(total, 60)
            from PIL import ImageDraw

            for index in range(limit):
                thumb = Image.new("RGB", (120, 160), (245, 243, 255))
                draw = ImageDraw.Draw(thumb)
                label = str(index + 1)
                draw.text((48, 68), label, fill=(124, 58, 237))
                buf = io.BytesIO()
                thumb.save(buf, format="PNG")
                encoded = base64.b64encode(buf.getvalue()).decode("ascii")
                pages.append(
                    {
                        "index": index,
                        "width": 120,
                        "height": 160,
                        "image": f"data:image/png;base64,{encoded}",
                    }
                )
            return JSONResponse({"pageCount": total, "truncated": total > limit, "pages": pages})
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Thumbnail error: {str(e)}")


# ==============================================================================
# 7. PDF TO EXCEL (Library: pdfplumber + pandas | pdfminer.six + openpyxl)
# ==============================================================================
@app.post("/api/pdf-to-excel")
async def pdf_to_excel(file: UploadFile = File(...)):
    try:
        content = await file.read()
        import pdfplumber
        import pandas as pd

        out_xlsx = os.path.join(TEMP_DIR, f"{os.path.splitext(file.filename)[0]}.xlsx")
        all_tables = []

        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                for table in tables:
                    if table:
                        df = pd.DataFrame(table[1:], columns=table[0])
                        all_tables.append(df)

        if all_tables:
            with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
                for idx, df in enumerate(all_tables):
                    df.to_excel(writer, sheet_name=f"Table_{idx+1}", index=False)
        else:
            df = pd.DataFrame([["No structured tables detected in document"]])
            df.to_excel(out_xlsx, index=False)

        return FileResponse(
            out_xlsx,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=f"{os.path.splitext(file.filename)[0]}.xlsx",
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF to Excel error: {str(e)}")


# ==============================================================================
# 8. OCR SCANNED PDF (Library: pytesseract + pdf2image | Tesseract Engine)
# ==============================================================================
@app.post("/api/ocr-pdf")
async def ocr_pdf(file: UploadFile = File(...)):
    try:
        content = await file.read()
        import pytesseract
        from pdf2image import convert_from_bytes

        images = convert_from_bytes(content)
        ocr_texts = []
        for img in images:
            text = pytesseract.image_to_string(img)
            ocr_texts.append(text)

        full_text = "\n\n--- Page Break ---\n\n".join(ocr_texts)
        out_txt = os.path.join(TEMP_DIR, "ocr_result.txt")
        with open(out_txt, "w", encoding="utf-8") as f:
            f.write(full_text)

        return FileResponse(out_txt, media_type="text/plain", filename="ocr_extracted.txt")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"OCR error: {str(e)}")


# ==============================================================================
# 9. WATERMARK / PAGE NUMBERS (Library: ReportLab + pypdf | Native Python)
# ==============================================================================
@app.post("/api/watermark-pdf")
async def watermark_pdf(
    file: UploadFile = File(...),
    watermark_text: str = Form("CONFIDENTIAL"),
    opacity: float = Form(0.25),
    mode: str = Form("text"),
    position: str = Form("center"),
    rotation: float = Form(45),
    scale: float = Form(0.28),
    logo: Optional[UploadFile] = File(None),
):
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.colors import Color
        from reportlab.lib.utils import ImageReader

        content = await file.read()
        reader = pypdf.PdfReader(io.BytesIO(content))
        writer = pypdf.PdfWriter()
        alpha = min(0.85, max(0.08, float(opacity)))
        angle = float(rotation)
        size_ratio = min(0.7, max(0.08, float(scale)))
        text = (watermark_text or "CONFIDENTIAL").strip()[:80]
        logo_bytes = await logo.read() if logo and logo.filename else None
        if mode == "logo" and not logo_bytes:
            raise HTTPException(status_code=400, detail="Upload a logo image to watermark with.")

        def _anchor(width: float, height: float):
            inset_x = width * 0.18
            inset_y = height * 0.18
            if position == "top-left":
                return inset_x, height - inset_y
            if position == "top-right":
                return width - inset_x, height - inset_y
            if position == "bottom-left":
                return inset_x, inset_y
            if position == "bottom-right":
                return width - inset_x, inset_y
            return width / 2, height / 2

        for page in reader.pages:
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)
            packet = io.BytesIO()
            can = canvas.Canvas(packet, pagesize=(width, height))
            x, y = _anchor(width, height)
            can.saveState()
            can.translate(x, y)
            can.rotate(angle)
            if mode == "logo" and logo_bytes:
                logo_img = Image.open(io.BytesIO(logo_bytes)).convert("RGBA")
                if alpha < 0.99:
                    r, g, b, a = logo_img.split()
                    logo_img = Image.merge("RGBA", (r, g, b, a.point(lambda px: int(px * alpha))))
                logo_buf = io.BytesIO()
                logo_img.save(logo_buf, format="PNG")
                logo_buf.seek(0)
                image = ImageReader(logo_buf)
                iw, ih = image.getSize()
                target_w = width * size_ratio
                target_h = target_w * (ih / max(1, iw))
                can.drawImage(image, -target_w / 2, -target_h / 2, width=target_w, height=target_h, mask="auto")
            else:
                can.setFillColor(Color(0.45, 0.45, 0.45, alpha=alpha))
                can.setFont("Helvetica-Bold", max(22, min(width, height) * size_ratio * 0.55))
                can.drawCentredString(0, 0, text)
            can.restoreState()
            can.save()
            packet.seek(0)
            page.merge_page(pypdf.PdfReader(packet).pages[0])
            writer.add_page(page)

        out_path = os.path.join(TEMP_DIR, "watermarked.pdf")
        with open(out_path, "wb") as f:
            writer.write(f)
        return FileResponse(out_path, media_type="application/pdf", filename="watermarked.pdf")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Watermark error: {str(e)}")


def _parse_pages(spec: str, page_count: int):
    raw = (spec or "").strip()
    if not raw or raw.lower() in ("all", "*"):
        return list(range(page_count))
    indices = []
    for part in raw.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            start, end = int(a.strip()), int(b.strip())
            if start > end:
                start, end = end, start
            indices.extend(range(max(1, start) - 1, min(page_count, end)))
        else:
            n = int(part)
            if 1 <= n <= page_count:
                indices.append(n - 1)
    seen = set()
    unique = []
    for i in indices:
        if 0 <= i < page_count and i not in seen:
            seen.add(i)
            unique.append(i)
    if not unique:
        raise HTTPException(status_code=400, detail="No valid pages matched that range.")
    return unique


@app.post("/api/split-pdf")
async def split_pdf(file: UploadFile = File(...), mode: str = Form("each"), pages: str = Form("all")):
    try:
        import zipfile

        content = await file.read()
        reader = pypdf.PdfReader(io.BytesIO(content))
        stem = os.path.splitext(file.filename or "document")[0]
        selected = _parse_pages(pages, len(reader.pages))
        # Split always returns a ZIP with one PDF per selected page.
        out_path = _temp_path(f"{stem}_pages.zip")
        with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for index in selected:
                writer = pypdf.PdfWriter()
                writer.add_page(reader.pages[index])
                buf = io.BytesIO()
                writer.write(buf)
                archive.writestr(f"{stem}_page_{index + 1}.pdf", buf.getvalue())
        return FileResponse(out_path, media_type="application/zip", filename=f"{stem}_pages.zip")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Split error: {str(e)}")


@app.post("/api/rotate-pdf")
async def rotate_pdf(file: UploadFile = File(...), angle: int = Form(90), pages: str = Form("all")):
    try:
        if angle not in (90, 180, 270):
            raise HTTPException(status_code=400, detail="Angle must be 90, 180, or 270.")
        reader = pypdf.PdfReader(io.BytesIO(await file.read()))
        rotate_set = set(_parse_pages(pages, len(reader.pages)))
        writer = pypdf.PdfWriter()
        for index, page in enumerate(reader.pages):
            if index in rotate_set:
                page.rotate(angle)
            writer.add_page(page)
        out_path = os.path.join(TEMP_DIR, "rotated.pdf")
        with open(out_path, "wb") as handle:
            writer.write(handle)
        return FileResponse(out_path, media_type="application/pdf", filename="rotated.pdf")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Rotate error: {str(e)}")


@app.post("/api/organize-pdf")
async def organize_pdf(
    file: UploadFile = File(...),
    mode: str = Form("keep"),
    pages: str = Form("1"),
    order: str = Form(""),
    rotations: str = Form(""),
):
    try:
        reader = pypdf.PdfReader(io.BytesIO(await file.read()))
        writer = pypdf.PdfWriter()
        page_count = len(reader.pages)

        if order.strip():
            sequence = []
            for part in order.split(","):
                part = part.strip()
                if not part:
                    continue
                index = int(part)
                if 0 <= index < page_count:
                    sequence.append(index)
            angles = []
            if rotations.strip():
                for part in rotations.split(","):
                    part = part.strip()
                    angles.append(int(part) if part else 0)
            if not sequence:
                raise HTTPException(status_code=400, detail="No pages left to save.")
            for offset, index in enumerate(sequence):
                page = reader.pages[index]
                extra = angles[offset] if offset < len(angles) else 0
                if extra in (90, 180, 270):
                    page.rotate(extra)
                writer.add_page(page)
        else:
            selected = set(_parse_pages(pages, page_count))
            for index, page in enumerate(reader.pages):
                keep = index in selected if mode == "keep" else index not in selected
                if keep:
                    writer.add_page(page)

        if len(writer.pages) == 0:
            raise HTTPException(status_code=400, detail="That would leave the PDF with zero pages.")
        out_path = os.path.join(TEMP_DIR, "organized.pdf")
        with open(out_path, "wb") as handle:
            writer.write(handle)
        return FileResponse(out_path, media_type="application/pdf", filename="organized.pdf")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Organize error: {str(e)}")


@app.post("/api/protect-pdf")
async def protect_pdf(file: UploadFile = File(...), password: str = Form(...)):
    try:
        if len(password) < 4:
            raise HTTPException(status_code=400, detail="Password must be at least 4 characters.")
        reader = pypdf.PdfReader(io.BytesIO(await file.read()))
        writer = pypdf.PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.encrypt(password, algorithm="AES-256")
        out_path = os.path.join(TEMP_DIR, "protected.pdf")
        with open(out_path, "wb") as handle:
            writer.write(handle)
        return FileResponse(out_path, media_type="application/pdf", filename="protected.pdf")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Protect error: {str(e)}")


@app.post("/api/unlock-pdf")
async def unlock_pdf(file: UploadFile = File(...), password: str = Form(...)):
    try:
        reader = pypdf.PdfReader(io.BytesIO(await file.read()))
        if reader.is_encrypted and reader.decrypt(password) == 0:
            raise HTTPException(status_code=400, detail="Incorrect PDF password.")
        writer = pypdf.PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        out_path = os.path.join(TEMP_DIR, "unlocked.pdf")
        with open(out_path, "wb") as handle:
            writer.write(handle)
        return FileResponse(out_path, media_type="application/pdf", filename="unlocked.pdf")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Unlock error: {str(e)}")


@app.post("/api/page-numbers")
async def page_numbers(
    file: UploadFile = File(...),
    position: str = Form("footer-center"),
    start: int = Form(1),
    number_format: str = Form("page-n-of-n"),
):
    try:
        from reportlab.pdfgen import canvas as rl_canvas

        reader = pypdf.PdfReader(io.BytesIO(await file.read()))
        writer = pypdf.PdfWriter()
        total = len(reader.pages)
        start_at = max(1, int(start))
        for index, page in enumerate(reader.pages):
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)
            number = start_at + index
            if number_format == "number":
                label = str(number)
            elif number_format == "page-n":
                label = f"Page {number}"
            else:
                label = f"Page {number} of {total}"
            packet = io.BytesIO()
            can = rl_canvas.Canvas(packet, pagesize=(width, height))
            can.setFillColorRGB(0.2, 0.2, 0.2)
            can.setFont("Helvetica", 11)
            y = height - 28 if position.startswith("header") else 22
            if position.endswith("left"):
                can.drawString(28, y, label)
            elif position.endswith("right"):
                can.drawRightString(width - 28, y, label)
            else:
                can.drawCentredString(width / 2, y, label)
            can.save()
            packet.seek(0)
            page.merge_page(pypdf.PdfReader(packet).pages[0])
            writer.add_page(page)
        out_path = os.path.join(TEMP_DIR, "numbered.pdf")
        with open(out_path, "wb") as handle:
            writer.write(handle)
        return FileResponse(out_path, media_type="application/pdf", filename="numbered.pdf")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Page numbers error: {str(e)}")


@app.post("/api/extract-text")
async def extract_text(file: UploadFile = File(...)):
    try:
        reader = pypdf.PdfReader(io.BytesIO(await file.read()))
        chunks = []
        for index, page in enumerate(reader.pages):
            text = (page.extract_text() or "").strip()
            chunks.append(f"--- Page {index + 1} ---\n{text if text else '[No digital text on this page]'}")
        out_path = os.path.join(TEMP_DIR, "extracted.txt")
        with open(out_path, "w", encoding="utf-8") as handle:
            handle.write("\n\n".join(chunks))
        return FileResponse(out_path, media_type="text/plain", filename="extracted.txt")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Extract text error: {str(e)}")


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    print(f"Starting TFRENZY Simple PDF Backend on http://0.0.0.0:{port} ...")
    uvicorn.run("backend.main:app", host="0.0.0.0", port=port, reload=False)

